#!/usr/bin/env python3
"""Review-only Paper 3 DOCX/PDF export; never submits or re-runs science."""
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile

import cairosvg

AREA = Path(__file__).resolve().parent
REPO = AREA.parents[3]
OUT = REPO / "artifacts/paper3_ijcip_review_only"

def run(args):
    subprocess.run(args, check=True)

def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def main():
    for exe in ("pandoc", "libreoffice", "pdftotext", "pdfinfo"):
        if shutil.which(exe) is None:
            raise RuntimeError("Missing export executable: " + exe)
    source = AREA / "PAPER3_IJCIP_MANUSCRIPT_R3.md"
    manuscript = source.read_text(encoding="utf-8")
    if manuscript.count("## Abstract\n") != 1 or not manuscript.startswith("# "):
        raise AssertionError("invalid source structure")
    body = manuscript.splitlines()[0] + "\n\n## Abstract\n" + manuscript.split("## Abstract\n",1)[1]
    OUT.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as directory:
        tmp = Path(directory)
        (tmp / "figures").mkdir()
        for name in ("PAPER3_IJCIP_FIG1_ARCHITECTURE", "PAPER3_IJCIP_FIG2_ENDPOINTS"):
            svg = AREA / "figures" / (name + ".svg")
            png = tmp / "figures" / (name + ".png")
            cairosvg.svg2png(url=str(svg), write_to=str(png), output_width=2400)
            shutil.copyfile(png, OUT / (name + "_REVIEW.png"))
            body = body.replace("figures/" + name + ".svg", "figures/" + name + ".png")
        markdown = tmp / "manuscript.md"
        markdown.write_text(body, encoding="utf-8")
        docx = OUT / "PAPER3_IJCIP_MANUSCRIPT_REVIEW_ONLY.docx"
        run(["pandoc","--from=gfm","--to=docx","--standalone","--resource-path",str(tmp),"-o",str(docx),str(markdown)])
        # Review layout only: render concise labels in tables, but preserve
        # the scientific manuscript text and frozen result identities.
        from docx import Document
        from docx.shared import Pt
        d = Document(docx)
        if len(d.tables) != 8:
            raise AssertionError("review DOCX expected eight original tables")
        aliases = {
            "D0_BASE":"D0", "L0_BASE":"L0",
            "D1_CORROBORATED":"D1", "L1_CORROBORATED":"L1"
        }
        for table_index, table in enumerate(d.tables):
            for row_index, row in enumerate(table.rows):
                for cell in row.cells:
                    if table_index in (4,5,6):
                        label = aliases.get(cell.text,cell.text)
                        if table_index==4 and row_index==0:
                            label = {
                                "Objective errors":"Errors",
                                "Unsafe proceed":"Unsafe",
                                "False-conservative hold":"FCH"
                            }.get(label,label)
                        if label != cell.text:
                            cell.text = label
                    for paragraph in cell.paragraphs:
                        for segment in paragraph.runs:
                            segment.font.size = Pt(8.5)
        # Preserve full scenario identifiers in this adjacent legend while
        # keeping the Study 7 example table readable at normal page width.
        s7_aliases = {
            "SAFE_CORROBORATED": "SAFE",
            "V5_INDEPENDENT_DISAGREEMENT": "V5 independent",
            "V5_CORRELATED_FALSE_CORROBORATION": "V5 correlated",
        }
        for row in d.tables[2].rows[1:]:
            name = row.cells[0].text
            if name in s7_aliases:
                row.cells[0].text = s7_aliases[name]
        s7_note = d.add_paragraph(
            "Study 7 scenario-label key: SAFE = SAFE_CORROBORATED; "
            "V5 independent = V5_INDEPENDENT_DISAGREEMENT; "
            "V5 correlated = V5_CORRELATED_FALSE_CORROBORATION."
        )
        for segment in s7_note.runs:
            segment.font.size=Pt(8.5)
        d.tables[2]._element.addnext(s7_note._element)
        note = d.add_paragraph(
            "Table-label key: D0 = D0_BASE; L0 = L0_BASE; "
            "D1 = D1_CORROBORATED; L1 = L1_CORROBORATED. "
            "Unsafe = unsafe proceed; FCH = false-conservative hold."
        )
        for segment in note.runs:
            segment.font.size=Pt(8.5)
        # Put the key BEFORE the table: otherwise the current PDF splits
        # policy rows and their legend across adjacent pages.
        d.tables[4]._element.addprevious(note._element)
        for paragraph in d.paragraphs:
            if paragraph.text.startswith("4.3 Held-out block results"):
                paragraph.paragraph_format.page_break_before=True
        d.save(docx)
        for filename,outname,fmt in (
            ("IJCIP_COVER_LETTER_AUTHOR_REVIEW_DRAFT_2026-10-20.md","IJCIP_COVER_LETTER_REVIEW_ONLY.docx","gfm"),
            ("IJCIP_HIGHLIGHTS_R1.txt","IJCIP_HIGHLIGHTS_REVIEW_ONLY.docx","markdown")
        ):
            run(["pandoc","-f",fmt,"-t","docx","--standalone","-o",str(OUT/outname),str(AREA/filename)])
        meta = (AREA/"IJCIP_SUBMISSION_METADATA_DRAFT.md").read_text(encoding="utf-8")
        title = tmp/"title.md"
        title.write_text("# Author details — author verification only\n\n"+meta+"\n",encoding="utf-8")
        run(["pandoc","-f","gfm","-t","docx","--standalone","-o",str(OUT/"IJCIP_TITLE_PAGE_REVIEW_ONLY.docx"),str(title)])
        run(["libreoffice","-env:UserInstallation="+(tmp/"lo_profile").as_uri(),"--headless","--convert-to","pdf","--outdir",str(OUT),str(docx)])
        pdf = OUT/"PAPER3_IJCIP_MANUSCRIPT_REVIEW_ONLY.pdf"
        if not pdf.is_file() or pdf.stat().st_size<5000:
            raise AssertionError("no PDF")
        extract=tmp/"check.txt"
        run(["pdftotext","-layout",str(pdf),str(extract)])
        readback=extract.read_text(encoding="utf-8")
        for term in ("196","784","F12","C0","References"):
            if term not in readback:
                raise AssertionError("PDF missing text: "+term)
        run(["pdfinfo",str(pdf)])
        receipt={"status":"REVIEW_ONLY_NOT_PUBLISHER_READY","source_sha256":digest(source),
                 "source_head":os.environ.get("PAPER3_SOURCE_HEAD","NOT_BOUND"),
                 "new_experiment":False,"journal_specific_rules_verified":False,
                 "author_confirmation":False,"visual_qa":False,"artifacts":{}}
        for artifact in sorted(OUT.iterdir()):
            if artifact.is_file() and artifact.name!="REVIEW_ONLY_SHA256_MANIFEST.json":
                receipt["artifacts"][artifact.name]={"sha256":digest(artifact),"bytes":artifact.stat().st_size}
        (OUT/"REVIEW_ONLY_SHA256_MANIFEST.json").write_text(json.dumps(receipt,indent=2)+"\n",encoding="utf-8")
    print("PAPER3_IJCIP_SOURCE_TO_WORD_AND_PDF=PASS_REVIEW_ONLY")
    print("PAPER3_IJCIP_SUBMISSION_AUTHORIZED=NO")

if __name__=="__main__":
    main()
