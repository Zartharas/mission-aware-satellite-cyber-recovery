#!/usr/bin/env python3
from __future__ import annotations

import json
import re
import shutil
import zipfile
from pathlib import Path

import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
import bibtexparser

ROOT = Path(__file__).resolve().parents[3]
PKG = ROOT / "publication" / "Paper_4_Study_8" / "IJSCCN"
LOCAL_PRIVATE = PKG / "_local_private"
LOCAL_SUBMISSION = PKG / "_local_submission"
PRIVATE_META_PATH = LOCAL_PRIVATE / "author_private.json"

MANUSCRIPT_SRC = PKG / "MANUSCRIPT_IJSCCN_R5.md"
BIB_SRC = PKG / "REFERENCES.bib"

FULL_TITLE = "Post-Quantum Trusted Recovery Under Intermittent Connectivity: Feasibility Across Modeled Contact Budgets and Public Observation-Opportunity Timing"
SHORT_TITLE = "Post-Quantum Satellite Recovery Under Intermittent Connectivity"

SAFE_AUTHOR_META = {
    "display_name": "AUTHOR_LOCAL_ONLY",
    "affiliation": "Independent Researcher",
    "location": "LOCATION_LOCAL_ONLY",
    "email": "EMAIL_LOCAL_ONLY",
    "orcid": "ORCID_LOCAL_ONLY",
    "telephone": "",
}

def load_author_metadata():
    if not PRIVATE_META_PATH.exists():
        return dict(SAFE_AUTHOR_META), False
    data = json.loads(PRIVATE_META_PATH.read_text(encoding="utf-8"))
    required = ("display_name", "affiliation", "location", "email", "orcid")
    missing = [key for key in required if not str(data.get(key, "")).strip()]
    if missing:
        raise SystemExit(
            "Private author metadata is incomplete: " + ", ".join(missing)
        )
    meta = dict(SAFE_AUTHOR_META)
    meta.update({key: str(value).strip() for key, value in data.items()})
    return meta, True

AUTHOR_META, PRIVATE_MODE = load_author_metadata()
OUT = (LOCAL_SUBMISSION / "build" / "r5") if PRIVATE_MODE else (PKG / "build" / "r5")
FIG = OUT / "figures"
OUT.mkdir(parents=True, exist_ok=True)
FIG.mkdir(parents=True, exist_ok=True)

AUTHOR = AUTHOR_META["display_name"]
AFFILIATION = ", ".join(
    part for part in (AUTHOR_META["affiliation"], AUTHOR_META["location"]) if part
)
EMAIL = AUTHOR_META["email"]
ORCID = AUTHOR_META["orcid"]
TELEPHONE = AUTHOR_META.get("telephone", "")

def render_author_tokens(text: str):
    telephone = AUTHOR_META.get("telephone", "").strip()
    if not telephone:
        text = re.sub(
            r"(?m)^Telephone:\s*\{\{AUTHOR_TELEPHONE\}\}\s*\n?",
            "",
            text,
        )

    replacements = {
        "{{AUTHOR_DISPLAY_NAME}}": AUTHOR_META["display_name"],
        "{{AUTHOR_AFFILIATION}}": AUTHOR_META["affiliation"],
        "{{AUTHOR_LOCATION}}": AUTHOR_META["location"],
        "{{AUTHOR_EMAIL}}": AUTHOR_META["email"],
        "{{AUTHOR_ORCID}}": AUTHOR_META["orcid"],
        "{{AUTHOR_TELEPHONE}}": telephone,
    }
    for token, value in replacements.items():
        text = text.replace(token, value)
    return text

def find_local_author_photo():
    for name in (
        "author_photo.png",
        "author_photo.jpg",
        "author_photo.jpeg",
        "author_photo.tif",
        "author_photo.tiff",
    ):
        candidate = LOCAL_PRIVATE / name
        if candidate.exists():
            return candidate
    return None

def set_style_font(style, font_name: str):
    style.font.name = font_name
    rfonts = style._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), font_name)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rfonts.attrib.pop(qn(f"w:{attr}"), None)

def set_run_font(run, font_name: str):
    run.font.name = font_name
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), font_name)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rfonts.attrib.pop(qn(f"w:{attr}"), None)

def apply_base_style(doc: Document):
    sec = doc.sections[0]
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    styles = doc.styles
    normal = styles["Normal"]
    set_style_font(normal, "Times New Roman")
    if "Default Paragraph Font" in styles:
        set_style_font(styles["Default Paragraph Font"], "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name in ["Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"]:
        st = styles[name]
        set_style_font(st, "Times New Roman")
    styles["Heading 1"].font.size = Pt(14)
    styles["Heading 2"].font.size = Pt(12)
    styles["Heading 3"].font.size = Pt(12)
    for name in ["Heading 1", "Heading 2", "Heading 3"]:
        styles[name].font.color.rgb = RGBColor(0, 0, 0)
        styles[name].font.bold = True
        styles[name].paragraph_format.keep_with_next = True
        styles[name].paragraph_format.space_before = Pt(10)
        styles[name].paragraph_format.space_after = Pt(4)
    styles["Heading 3"].font.bold = False
    styles["Heading 3"].font.italic = True

def add_formatted_runs(p, text: str):
    tick = "\x60"
    token_re = r"(\*\*.*?\*\*|" + tick + r".*?" + tick + r"|\*.*?\*)"
    parts = re.split(token_re, text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            r.bold = True
        elif part.startswith(tick) and part.endswith(tick):
            r = p.add_run(part[1:-1])
            r.font.name = "Courier New"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
        elif part.startswith("*") and part.endswith("*"):
            r = p.add_run(part[1:-1])
            r.italic = True
        else:
            p.add_run(part)


def _table_cells(line: str):
    return [cell.strip() for cell in line.strip().strip("|").split("|")]

def _is_table_separator(line: str):
    cells = _table_cells(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells)

def add_markdown_table(doc: Document, rows):
    if not rows:
        return
    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    header = [cell.replace("**", "").strip() for cell in normalized[0]]

    if width == 3 and header[:3] == ["Factor", "Levels", "Frozen values / interpretation"]:
        column_widths = [Inches(1.45), Inches(0.65), Inches(4.40)]
    elif width == 2 and header[:2] == ["Quantity", "Frozen value"]:
        column_widths = [Inches(3.25), Inches(3.25)]
    else:
        column_widths = [Inches(6.5 / width)] * width

    table = doc.add_table(rows=len(normalized), cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for j, column_width in enumerate(column_widths):
        table.columns[j].width = column_width

    for i, row in enumerate(normalized):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.width = column_widths[j]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_together = True
            if i == 0:
                run = p.add_run(value.replace("**", ""))
                run.bold = True
                set_run_font(run, "Times New Roman")
            else:
                add_formatted_runs(p, value)
    doc.add_paragraph()

def render_markdown_lines(doc: Document, lines):
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped or stripped == "---" or stripped.startswith("<!--"):
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and _is_table_separator(lines[i + 1]):
            rows = [_table_cells(line)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(_table_cells(lines[i]))
                i += 1
            add_markdown_table(doc, rows)
            continue
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            heading = line[level:].strip()
            if heading.startswith("Table 1.") or heading.startswith("Table 2."):
                doc.add_page_break()
            doc.add_heading(heading, level=min(max(level - 1, 1), 3))
            i += 1
            continue
        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.35)
            r = p.add_run(line[2:].replace("**", ""))
            r.italic = True
            i += 1
            continue
        if re.match(r"^\d+\.\s+", line):
            p = doc.add_paragraph(style="List Number")
            add_formatted_runs(p, re.sub(r"^\d+\.\s+", "", line))
            i += 1
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_formatted_runs(p, line[2:])
            i += 1
            continue
        p = doc.add_paragraph()
        add_formatted_runs(p, line)
        i += 1

def assert_docx_integrity(path: Path):
    with zipfile.ZipFile(path, "r") as z:
        names = set(z.namelist())
        if "word/comments.xml" in names:
            raise SystemExit(f"Comments part unexpectedly present in {path}")
        xml_names = [
            name for name in names
            if name.startswith("word/") and name.endswith(".xml")
        ]
        xml = "\n".join(
            z.read(name).decode("utf-8", errors="ignore") for name in xml_names
        )
    if "—" in xml:
        raise SystemExit(f"Em dash found in publisher-facing DOCX XML: {path}")
    markers = {
        "inserted revision": r"<w:ins(?:\\s|>)",
        "deleted revision": r"<w:del(?:\\s|>)",
        "hidden text": r"<w:vanish(?:\\s|/|>)",
    }
    for label, pattern in markers.items():
        if re.search(pattern, xml):
            raise SystemExit(f"{label} marker found in {path}")

def load_bib():
    with BIB_SRC.open("r", encoding="utf-8") as f:
        return bibtexparser.load(f)

def citation_order(text: str):
    order = []
    seen = set()
    pat = re.compile(r"\[([^\]]*@[^]]+)\]")
    for m in pat.finditer(text):
        inside = m.group(1)
        for key in re.findall(r"@([A-Za-z0-9_:-]+)", inside):
            if key not in seen:
                seen.add(key)
                order.append(key)
    return order

def replace_citations(text: str, order):
    index = {k:i+1 for i,k in enumerate(order)}
    def repl(m):
        inside = m.group(1)
        keys = re.findall(r"@([A-Za-z0-9_:-]+)", inside)
        if not keys:
            return m.group(0)
        nums = [str(index[k]) for k in keys if k in index]
        return "[" + ", ".join(nums) + "]"
    return re.sub(r"\[([^\]]*@[^]]+)\]", repl, text)

def clean_latex_text(raw: str):
    if not raw:
        return ""
    replacements = {
        '{\"a}': 'ä',
        '{\"u}': 'ü',
        '{\"o}': 'ö',
        '{\\\'e}': 'é',
        '{\\\'a}': 'á',
        '{\\\'i}': 'í',
        '{\\\'o}': 'ó',
        '{\\\'u}': 'ú',
        '---': ' - ',
        '--': '-',
        '—': ' - ',
        '–': '-',
    }
    text = raw
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = text.replace('{', '').replace('}', '')
    text = text.replace('\\&', '&')
    return text

def normalize_author_names(raw: str):
    if not raw:
        return ""
    cleaned = clean_latex_text(raw)
    if "National Institute of Standards" in cleaned and "Technology" in cleaned:
        return "National Institute of Standards and Technology"
    parts = [clean_latex_text(p.strip()) for p in raw.replace("\n", " ").split(" and ")]
    return ", ".join(parts)

def format_reference(entry):
    authors = normalize_author_names(entry.get("author") or entry.get("organization") or "")
    title = clean_latex_text(entry.get("title",""))
    journal = clean_latex_text(
        entry.get("journal")
        or entry.get("booktitle")
        or entry.get("institution")
        or entry.get("howpublished")
        or ""
    )
    year = entry.get("year","")
    volume = entry.get("volume","")
    number = entry.get("number","")
    pages = entry.get("pages","").replace("--","-")
    doi = entry.get("doi","")
    url = entry.get("url","")
    note = clean_latex_text(entry.get("note",""))
    std_number = clean_latex_text(entry.get("number",""))
    pieces = []
    if authors: pieces.append(authors.rstrip(".") + ".")
    if title: pieces.append(title.rstrip(".") + ".")
    tail = ""
    if journal: tail += journal
    if std_number and not volume and std_number not in tail:
        tail += (" " if tail else "") + std_number
    if year: tail += (" " if tail else "") + year
    if volume:
        tail += f"; {volume}"
        if number: tail += f"({number})"
    if pages: tail += f": {pages}"
    if tail: pieces.append(tail.rstrip(".") + ".")
    if doi:
        pieces.append("DOI: " + doi + ".")
    elif url:
        access = " Accessed 2026-09-21." if "access" not in note.lower() else " " + note.rstrip(".") + "."
        pieces.append(url + "." + access)
    return " ".join(pieces)

def make_figures():
    states = [
        "COMPROMISED",
        "RECOVERY AUTHORITY\nESTABLISHED",
        "SUCCESSOR PROFILE\nSELECTED",
        "SUCCESSOR KEY\nMATERIAL STAGED",
        "TRANSITION PROOF\nACCEPTED",
        "NEW EPOCH\nCOMMITTED",
        "OLD EPOCH\nREVOKED",
        "TRUST\nRESTORED",
    ]
    fig, ax = plt.subplots(figsize=(9.2, 5.2))
    ax.axis("off")
    coords = [
        (0.08,0.72),(0.36,0.72),(0.64,0.72),(0.92,0.72),
        (0.92,0.28),(0.64,0.28),(0.36,0.28),(0.08,0.28),
    ]
    for i,((x,y),s) in enumerate(zip(coords,states)):
        ax.text(x, y, s, ha="center", va="center", fontsize=9.5,
                bbox=dict(boxstyle="round,pad=0.40", fill=False))
        if i < len(states)-1:
            nx,ny=coords[i+1]
            if abs(ny-y) < 0.01:
                direction = 1 if nx > x else -1
                ax.annotate("", xy=(nx-0.09*direction,y), xytext=(x+0.09*direction,y),
                            arrowprops=dict(arrowstyle="->"))
            else:
                ax.annotate("", xy=(nx,ny+0.09), xytext=(x,y-0.09),
                            arrowprops=dict(arrowstyle="->"))
    ax.set_xlim(-0.04,1.04); ax.set_ylim(0,1)
    fig.tight_layout()
    fig.savefig(FIG/"Figure_1_Recovery_State_Machine.tiff", dpi=800, bbox_inches="tight")
    fig.savefig(FIG/"Figure_1_Recovery_State_Machine.png", dpi=300, bbox_inches="tight")
    plt.close(fig)

    labels = ["ML-KEM-512 /\nML-DSA-44","ML-KEM-768 /\nML-DSA-65","ML-KEM-1024 /\nML-DSA-87"]
    vals = [93.7500,64.9306,61.8056]
    fig, ax = plt.subplots(figsize=(6.8,4.2))
    bars = ax.bar(labels, vals)
    ax.set_ylabel("Modeled trusted-recovery success (%)")
    ax.set_ylim(0,100)
    ax.set_title("Study 8: fixed-capacity success by cryptographic profile")
    for b,v in zip(bars,vals):
        ax.text(b.get_x()+b.get_width()/2, v+2, f"{v:.4f}%", ha="center", va="bottom", fontsize=9)
    fig.tight_layout()
    fig.savefig(FIG/"Figure_2_Study8_Profile_Success.tiff", dpi=800, bbox_inches="tight")
    fig.savefig(FIG/"Figure_2_Study8_Profile_Success.png", dpi=300, bbox_inches="tight")
    plt.close(fig)

    labels = ["6 h","12 h","24 h"]
    vals = [5.2863,20.5947,55.0661]
    fig, ax = plt.subplots(figsize=(6.8,4.2))
    bars = ax.bar(labels, vals)
    ax.set_ylabel("Cases with finite modeled rate threshold (%)")
    ax.set_ylim(0,60)
    ax.set_title("Study 8E: finite threshold fraction by elapsed-time horizon")
    for b,v in zip(bars,vals):
        ax.text(b.get_x()+b.get_width()/2, v+1.5, f"{v:.4f}%", ha="center", va="bottom", fontsize=9)
    fig.tight_layout()
    fig.savefig(FIG/"Figure_3_Study8E_Horizon_Finite_Fraction.tiff", dpi=800, bbox_inches="tight")
    fig.savefig(FIG/"Figure_3_Study8E_Horizon_Finite_Fraction.png", dpi=300, bbox_inches="tight")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(10,4.8))
    ax.axis("off")
    ax.text(0.5,0.89,"Post-compromise post-quantum satellite recovery",
            ha="center",va="center",fontsize=15,weight="bold")
    ax.text(0.24,0.62,"STUDY 8\nFixed modeled contact capacity\nDoes recovery fit?",
            ha="center",va="center",fontsize=11,
            bbox=dict(boxstyle="round,pad=0.62",fill=False))
    ax.text(0.76,0.62,"STUDY 8E\nPublic observation-opportunity timing\nWhat modeled rate is required?",
            ha="center",va="center",fontsize=11,
            bbox=dict(boxstyle="round,pad=0.62",fill=False))
    ax.text(0.50,0.62,"SEPARATE EVIDENCE\nNo pooling\nNo slot-to-seconds conversion",
            ha="center",va="center",fontsize=8.8,weight="bold")
    ax.text(0.5,0.34,
            "P3 versus P1: no feasibility advantage in either frozen evidence layer",
            ha="center",va="center",fontsize=10.5)
    ax.text(0.5,0.24,
            "Object burden + opportunity timing + recovery horizon constrain modeled feasibility",
            ha="center",va="center",fontsize=10)
    ax.text(0.5,0.08,
            "SatNOGS supplies timing proxies only - not authenticated command contacts or measured throughput",
            ha="center",va="center",fontsize=8)
    fig.tight_layout()
    fig.savefig(FIG/"IJSCCN_GTOC.tiff", dpi=800, bbox_inches="tight")
    fig.savefig(FIG/"IJSCCN_GTOC.png", dpi=300, bbox_inches="tight")
    fig.savefig(FIG/"IJSCCN_GTOC.svg", bbox_inches="tight")
    plt.close(fig)

def add_title_page(doc):
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(FULL_TITLE); r.bold=True; r.font.size=Pt(16); set_run_font(r, "Times New Roman")
    p.paragraph_format.space_after=Pt(10)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(f"{AUTHOR}, MS, PhD"); r.bold=True; set_run_font(r, "Times New Roman")
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(AFFILIATION)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"Corresponding author: {EMAIL}")
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"ORCID: {ORCID}")
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"Short title: {SHORT_TITLE}")
    p.paragraph_format.space_after=Pt(12)

def insert_figure(doc, png_path, caption):
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(png_path), width=Inches(6.4))
    cp=doc.add_paragraph()
    cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    rr=cp.add_run(caption); rr.italic=True

def render_markdown_to_docx(md_text, output_docx, bib):
    md_text = render_author_tokens(md_text)
    order = citation_order(md_text)
    md_text = replace_citations(md_text, order)
    marker = "<!-- R5_REFERENCES_INSERTION_POINT -->"
    if marker not in md_text:
        raise SystemExit("R5 reference insertion marker missing")
    body_text, post_reference_text = md_text.split(marker, 1)

    lines = body_text.splitlines()
    abstract_idx = next(
        (i for i, line in enumerate(lines) if line.strip() == "## Abstract"),
        None,
    )
    if abstract_idx is None:
        raise SystemExit("R5 Abstract heading missing")

    doc = Document()
    apply_base_style(doc)
    doc.core_properties.title = FULL_TITLE
    doc.core_properties.subject = "IJSCCN Original Paper submission manuscript"
    doc.core_properties.author = AUTHOR
    add_title_page(doc)

    render_markdown_lines(doc, lines[abstract_idx:])

    doc.add_heading("References", level=1)
    entries = {e["ID"]: e for e in bib.entries}
    missing = [key for key in order if key not in entries]
    if missing:
        raise SystemExit("Missing bibliography entries: " + ", ".join(missing))
    for i, key in enumerate(order, 1):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(4)
        p.add_run(f"{i}. ")
        p.add_run(format_reference(entries[key]))

    render_markdown_lines(doc, post_reference_text.splitlines())

    doc.save(output_docx)
    assert_docx_integrity(output_docx)
    return order

def markdown_to_cover_letter_docx(md_path, out_path):
    doc=Document()
    apply_base_style(doc)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("Cover Letter"); r.bold=True; r.font.size=Pt(15)

    rendered = render_author_tokens(md_path.read_text(encoding="utf-8"))
    lines = rendered.splitlines()
    i = 0

    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip() or line.startswith("# "):
            i += 1
            continue

        if line.strip() == "Sincerely,":
            p=doc.add_paragraph()
            p.paragraph_format.space_after=Pt(2)
            add_formatted_runs(p, "Sincerely,")
            i += 1
            while i < len(lines) and not lines[i].strip():
                i += 1

            signature = []
            while i < len(lines) and lines[i].strip():
                signature.append(lines[i].rstrip())
                i += 1

            p=doc.add_paragraph()
            p.paragraph_format.space_after=Pt(0)
            for index, signature_line in enumerate(signature):
                if index:
                    p.add_run().add_break()
                add_formatted_runs(p, signature_line)
            continue

        if line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("- "):
            p=doc.add_paragraph(style="List Bullet")
            add_formatted_runs(p,line[2:])
        else:
            p=doc.add_paragraph()
            add_formatted_runs(p,line)
        i += 1

    doc.save(out_path)
    assert_docx_integrity(out_path)

def markdown_to_simple_docx(md_path, out_path, title):
    doc=Document()
    apply_base_style(doc)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(title); r.bold=True; r.font.size=Pt(15)
    rendered = render_author_tokens(md_path.read_text(encoding="utf-8"))
    for line in rendered.splitlines():
        if not line.strip() or line.startswith("# "):
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("- "):
            p=doc.add_paragraph(style="List Bullet"); add_formatted_runs(p,line[2:])
        else:
            p=doc.add_paragraph(); add_formatted_runs(p,line)
    doc.save(out_path)
    assert_docx_integrity(out_path)

def main():
    make_figures()
    bib=load_bib()
    md=MANUSCRIPT_SRC.read_text(encoding="utf-8")
    order=render_markdown_to_docx(md, OUT/"MANUSCRIPT_IJSCCN_R5.docx", bib)

    markdown_to_cover_letter_docx(PKG/"COVER_LETTER.md", OUT/"COVER_LETTER_IJSCCN_R5.docx")
    markdown_to_simple_docx(PKG/"AUTHOR_BIOGRAPHY.md", OUT/"AUTHOR_BIOGRAPHY_R5.docx", "Author Biography")
    markdown_to_simple_docx(PKG/"TITLE_PAGE.md", OUT/"TITLE_PAGE_IJSCCN_R5.docx", "Title Page")
    markdown_to_simple_docx(PKG/"AI_USE_DECLARATION.md", OUT/"AI_USE_DECLARATION_R5.docx", "AI Use Declaration")

    photo = find_local_author_photo()
    if PRIVATE_MODE and photo is not None:
        shutil.copy2(photo, OUT / ("AUTHOR_PHOTOGRAPH" + photo.suffix.lower()))

    manifest={
        "package_id":"P4-IJSCCN-PACKAGE-R5-QA-001",
        "source_manuscript":str(MANUSCRIPT_SRC.relative_to(ROOT)),
        "reference_count":len(order),
        "figures":[p.name for p in sorted(FIG.iterdir())],
        "scientific_reanalysis_performed":False,
        "publisher_submission_authorized":False,
        "private_submission_overlay_used":PRIVATE_MODE,
        "author_photo_included":bool(PRIVATE_MODE and photo is not None),
        "output_scope":"LOCAL_PRIVATE_SUBMISSION" if PRIVATE_MODE else "PUBLIC_SAFE_QA",
        "r5_source_preservation_audit":"PASS",
        "tables_after_references":True,
        "scientific_figures_separate":True,
        "zero_em_dash_enforced":True,
        "tracked_changes_comments_hidden_text_enforced_absent":True
    }
    (OUT/"BUILD_MANIFEST.json").write_text(json.dumps(manifest,indent=2)+"\n",encoding="utf-8")

    zip_path=OUT/"Paper4_IJSCCN_Submission_Package_R5.zip"
    with zipfile.ZipFile(zip_path,"w",zipfile.ZIP_DEFLATED) as z:
        for p in sorted(OUT.rglob("*")):
            if p==zip_path or p.is_dir():
                continue
            z.write(p,p.relative_to(OUT))
        for source in [
            PKG/"GTOC.md",
            PKG/"DATA_AVAILABILITY_STATEMENT.md",
            PKG/"SUBMISSION_CHECKLIST.md",
            PKG/"R5_WILEY_IJSCCN_EVIDENCE_AUDIT_2026-09-22.md",
            PKG/"R5_LIVE_POLICY_REQUIREMENT_MATRIX_2026-09-21.md",
            PKG/"R5_SOURCE_PRESERVATION_AUDIT_2026-09-22.md",
            PKG/"PACKAGE_STATUS.json",
        ]:
            z.write(source,source.name)
    print(json.dumps(manifest,indent=2))
    print(f"ZIP={zip_path}")

if __name__=="__main__":
    main()
